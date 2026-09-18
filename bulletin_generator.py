from __future__ import annotations

import html
import http.cookiejar
import re
import sys
import time
import urllib.error
import urllib.request
import zipfile
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin, urlparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook


DEFAULT_SECTIONS = [
    "SBP Related News / Press Release",
    "Domestic News",
    "Editorials / Opinion / Analysis",
    "Foreign Media Updates",
]

HEADER_ALIASES = {
    "include": ["include", "use", "selected"],
    "section": ["section", "category", "domain"],
    "story_group": ["storygroup", "storytopic", "story", "topic"],
    "headline": ["headline", "articleheadline", "newstitle", "title", "heading"],
    "source": ["source", "newspaper", "publication", "paper"],
    "url": ["url", "link", "articlelink"],
    "article_text": ["articletext", "manualtext", "content", "body"],
    "author": ["author", "byline", "writer"],
    "date_text": ["date", "published", "publishdate", "articledate"],
    "language": ["language", "lang"],
    "priority": ["priority"],
    "notes": ["notes", "printreference", "reference"],
    "status": ["status", "extractionstatus"],
}

BLUE = RGBColor(31, 111, 104)
DARK_BLUE = RGBColor(20, 61, 58)
NAVY = RGBColor(23, 50, 77)
GRAY = RGBColor(83, 93, 104)

BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/140.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Cache-Control": "no-cache",
    "Upgrade-Insecure-Requests": "1",
}

_BUSINESS_RECORDER_CANDIDATES: list["SyndicatedCandidate"] | None = None
_DAWN_CANDIDATES: list["SyndicatedCandidate"] | None = None
_AIB_CANDIDATES: list["SyndicatedCandidate"] | None = None
_STREETINSIDER_CANDIDATES: list["SyndicatedCandidate"] | None = None


@dataclass
class NewsItem:
    row_number: int
    section: str
    story_group: str
    headline: str
    source: str
    author: str
    url: str
    date_text: str
    article_text: str
    language: str
    priority: str
    notes: str
    status: str


@dataclass(frozen=True)
class SyndicatedCandidate:
    title: str
    url: str
    source: str


class ParagraphCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_script = False
        self.in_style = False
        self.in_title = False
        self.in_heading = False
        self.in_paragraph = False
        self.title_parts: list[str] = []
        self.heading_parts: list[str] = []
        self.headings: list[str] = []
        self.current: list[str] = []
        self.paragraphs: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag == "script":
            self.in_script = True
        elif tag == "style":
            self.in_style = True
        elif tag == "title":
            self.in_title = True
            self.title_parts = []
        elif tag in {"h1", "h2"}:
            self.in_heading = True
            self.heading_parts = []
        elif tag == "p":
            self.in_paragraph = True
            self.current = []

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "script":
            self.in_script = False
        elif tag == "style":
            self.in_style = False
        elif tag == "title":
            self.in_title = False
        elif tag in {"h1", "h2"} and self.in_heading:
            text = clean_text(" ".join(self.heading_parts))
            if len(text) >= 12:
                self.headings.append(text)
            self.in_heading = False
            self.heading_parts = []
        elif tag == "p" and self.in_paragraph:
            text = clean_text(" ".join(self.current))
            if len(text) >= 45:
                self.paragraphs.append(text)
            self.in_paragraph = False
            self.current = []

    def handle_data(self, data: str) -> None:
        if self.in_script or self.in_style:
            return
        if self.in_title:
            self.title_parts.append(data)
        if self.in_heading:
            self.heading_parts.append(data)
        if self.in_paragraph:
            self.current.append(data)

    def best_title(self) -> str:
        for candidate in self.headings:
            if candidate:
                return candidate
        title = clean_text(" ".join(self.title_parts))
        title = re.split(r"\s+[|-]\s+", title)[0].strip()
        return title


class AnchorCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.current_attrs: dict[str, str] | None = None
        self.current_text: list[str] = []
        self.anchors: list[tuple[dict[str, str], str]] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() == "a" and self.current_attrs is None:
            self.current_attrs = {str(key).lower(): str(value or "") for key, value in attrs}
            self.current_text = []

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self.current_attrs is not None:
            self.anchors.append((self.current_attrs, clean_text(" ".join(self.current_text))))
            self.current_attrs = None
            self.current_text = []

    def handle_data(self, data: str) -> None:
        if self.current_attrs is not None:
            self.current_text.append(data)


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    text = html.unescape(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def first_alias_match(header_map: dict[str, int], field: str) -> int | None:
    for alias in HEADER_ALIASES[field]:
        if alias in header_map:
            return header_map[alias]
    return None


def is_http_url(value: str) -> bool:
    parsed = urlparse(value or "")
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def requires_ocr(value: str) -> bool:
    parsed = urlparse(value or "")
    path = parsed.path.lower()
    query = parsed.query.lower()
    extensions = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".pdf")
    if path.endswith(extensions):
        return True
    return any(ext in query for ext in extensions) and any(
        key in query for key in ("newssrc=", "image=", "img=")
    )


def normalize_section_label(value: str) -> str:
    text = clean_text(value)
    upper = text.upper()
    if "SBP" in upper:
        return "SBP Related News / Press Release"
    if "DOMESTIC" in upper:
        return "Domestic News"
    if "EDITORIAL" in upper or "OPINION" in upper or "ARTICLE" in upper:
        return "Editorials / Opinion / Analysis"
    if "FOREIGN" in upper or "INTERNATIONAL" in upper:
        return "Foreign Media Updates"
    return text or "Domestic News"


def source_meta(item: "NewsItem") -> str:
    parts = [part for part in [item.source, item.author, item.date_text] if part]
    return " / ".join(parts) or "Unspecified source"


def find_header_row(sheet) -> tuple[int, dict[str, int]]:
    for row in range(1, min(sheet.max_row, 12) + 1):
        normalized = {
            normalize_header(sheet.cell(row=row, column=col).value): col
            for col in range(1, sheet.max_column + 1)
            if clean_text(sheet.cell(row=row, column=col).value)
        }
        matches = {
            field: first_alias_match(normalized, field)
            for field in HEADER_ALIASES
        }
        required_found = [matches["section"], matches["headline"], matches["source"]]
        if all(required_found):
            return row, {key: value for key, value in matches.items() if value is not None}
    raise ValueError("Could not find the input header row. Please use the official empty template.")


def read_news_items(excel_path: Path) -> tuple[list[NewsItem], list[str]]:
    if getattr(extract_article_from_url, "__module__", "") == __name__:
        reset_source_caches()
    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook["News Input"] if "News Input" in workbook.sheetnames else workbook.active
    header_row, columns = find_header_row(sheet)
    warnings: list[str] = []
    items: list[NewsItem] = []
    last_section = ""
    extraction_results: dict[tuple[str, str, str], tuple[str, str, str]] = {}

    def cell_text(row: int, field: str) -> str:
        col = columns.get(field)
        if not col:
            return ""
        return clean_text(sheet.cell(row=row, column=col).value)

    extraction_requests: dict[tuple[str, str, str], None] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        include = cell_text(row, "include").lower()
        if include in {"no", "n", "false", "0", "skip", "skipped"}:
            continue
        article_text = cell_text(row, "article_text")
        url = cell_text(row, "url")
        if article_text or not is_http_url(url):
            continue
        headline = cell_text(row, "headline") or cell_text(row, "story_group")
        source = cell_text(row, "source") or "Unspecified source"
        extraction_requests[(url, headline, source)] = None

    is_default_extractor = getattr(extract_article_from_url, "__module__", "") == __name__
    has_reuters_request = any(
        (urlparse(url).hostname or "").lower().endswith("reuters.com")
        or source.lower() == "reuters"
        for url, _, source in extraction_requests
    )
    if is_default_extractor and has_reuters_request:
        with ThreadPoolExecutor(max_workers=3) as pool:
            preload = [
                pool.submit(provider)
                for provider in (
                    business_recorder_candidates,
                    aib_candidates,
                    streetinsider_candidates,
                )
            ]
            for future in as_completed(preload):
                try:
                    future.result()
                except Exception:
                    pass

    if extraction_requests:
        with ThreadPoolExecutor(max_workers=min(8, len(extraction_requests))) as pool:
            pending = {
                pool.submit(
                    extract_article_from_url,
                    url,
                    headline,
                    source,
                ): (url, headline, source)
                for url, headline, source in extraction_requests
            }
            for future in as_completed(pending):
                key = pending[future]
                try:
                    extraction_results[key] = future.result()
                except Exception as exc:
                    extraction_results[key] = (
                        "[The app could not extract this article automatically. Please paste the article text in the Excel template.]",
                        f"URL extraction failed for {key[0]}: {exc}",
                        "",
                    )

    for row in range(header_row + 1, sheet.max_row + 1):
        raw_values = [clean_text(sheet.cell(row=row, column=col).value) for col in range(1, sheet.max_column + 1)]
        if not any(raw_values):
            continue

        include = cell_text(row, "include").lower()
        if include in {"no", "n", "false", "0", "skip", "skipped"}:
            continue

        raw_section = cell_text(row, "section")
        if raw_section:
            last_section = raw_section
        section = normalize_section_label(raw_section or last_section)
        headline = cell_text(row, "headline")
        source = cell_text(row, "source") or "Unspecified source"
        story_group = cell_text(row, "story_group") or headline
        url = cell_text(row, "url")
        author = cell_text(row, "author")
        date_text = cell_text(row, "date_text")
        article_text = cell_text(row, "article_text")
        language = cell_text(row, "language") or "English"
        priority = cell_text(row, "priority")
        notes = cell_text(row, "notes")
        status = cell_text(row, "status")

        if not headline and story_group:
            headline = story_group

        status_lower = status.lower()
        if not article_text and is_http_url(url):
            article_text, extraction_note, extracted_title = extraction_results[
                (url, headline, source)
            ]
            if not headline and extracted_title:
                headline = extracted_title
            if extraction_note:
                if status_lower:
                    extraction_note = f"{extraction_note} Uploaded status: {status}."
                warnings.append(f"Row {row}: {extraction_note}")
            elif status_lower and not status_lower.startswith("ok") and not status_lower.startswith("partial"):
                status = "recovered"
        elif not article_text and url and not is_http_url(url):
            recovered = recover_article_without_url(headline, source)
            if recovered:
                article_text, recovered_title, recovered_url = recovered
                url = recovered_url
                if not headline and recovered_title:
                    headline = recovered_title
                status = "recovered"
            else:
                article_text = "[Article text was not provided and the Link cell is not a reachable web URL. Please paste the article text or a valid URL before final circulation.]"
                warnings.append(f"Row {row}: Link is not a valid URL: {url}")
        elif not article_text:
            article_text = "[Article text was not provided. Add article text or a reachable URL before final circulation.]"
            warnings.append(f"Row {row}: no article text or URL was provided.")

        if not headline:
            headline = f"{source} article"
            warnings.append(f"Row {row}: Headline is empty; used a placeholder headline.")

        if not story_group:
            story_group = headline

        items.append(
            NewsItem(
                row_number=row,
                section=section,
                story_group=story_group or headline,
                headline=headline,
                source=source,
                author=author,
                url=url,
                date_text=date_text,
                article_text=article_text,
                language=language,
                priority=priority,
                notes=notes,
                status=status,
            )
        )

    if not items:
        raise ValueError("No included news rows were found. Add at least one row with Include = Yes.")
    return items, warnings


def fetch_markup(
    url: str,
    *,
    timeout: int = 15,
    opener=None,
) -> str:
    request = urllib.request.Request(url, headers=BROWSER_HEADERS)
    open_request = opener.open if opener is not None else urllib.request.urlopen
    for attempt in range(2):
        try:
            with open_request(request, timeout=timeout) as response:
                media_type = (response.headers.get_content_type() or "").lower()
                if media_type.startswith("image/") or media_type == "application/pdf":
                    raise ValueError("the URL points to an image or PDF and requires OCR")
                charset = response.headers.get_content_charset() or "utf-8"
                raw = response.read(1_500_000)
            return raw.decode(charset, errors="replace")
        except urllib.error.HTTPError as exc:
            if attempt == 1 or exc.code not in {429, 500, 502, 503, 504}:
                raise
        except (urllib.error.URLError, TimeoutError):
            if attempt == 1:
                raise
        time.sleep(0.35 * (attempt + 1))
    raise RuntimeError(f"Failed to fetch {url}")


def extract_from_markup(markup: str) -> tuple[str, str]:
    collector = ParagraphCollector()
    collector.feed(markup)
    extracted_title = collector.best_title()
    paragraphs = [paragraph for paragraph in collector.paragraphs if not is_boilerplate(paragraph)]
    if paragraphs:
        return "\n\n".join(paragraphs[:24]), extracted_title

    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", markup)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return clean_text(text)[:6000], extracted_title


def normalize_headline(value: str) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"^(?:rpt-|update\s*\d*[-:]|forex-|analysis-)+\s*", "", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def headline_score(expected: str, candidate: str) -> float:
    left = normalize_headline(expected)
    right = normalize_headline(candidate)
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    left_tokens = set(left.split())
    right_tokens = set(right.split())
    token_score = len(left_tokens & right_tokens) / max(1, len(left_tokens | right_tokens))
    sequence_score = SequenceMatcher(None, left, right).ratio()
    return max(token_score, sequence_score)


def collect_anchors(markup: str) -> list[tuple[dict[str, str], str]]:
    collector = AnchorCollector()
    collector.feed(markup)
    return collector.anchors


def reset_source_caches() -> None:
    global _BUSINESS_RECORDER_CANDIDATES
    global _DAWN_CANDIDATES
    global _AIB_CANDIDATES
    global _STREETINSIDER_CANDIDATES
    _BUSINESS_RECORDER_CANDIDATES = None
    _DAWN_CANDIDATES = None
    _AIB_CANDIDATES = None
    _STREETINSIDER_CANDIDATES = None


def business_recorder_candidates() -> list[SyndicatedCandidate]:
    global _BUSINESS_RECORDER_CANDIDATES
    if _BUSINESS_RECORDER_CANDIDATES is not None:
        return _BUSINESS_RECORDER_CANDIDATES

    candidates: list[SyndicatedCandidate] = []
    try:
        markup = fetch_markup("https://www.brecorder.com/latest-news/", timeout=18)
        for attrs, title in collect_anchors(markup):
            href = attrs.get("href", "")
            if len(title) < 20 or "/news/" not in href:
                continue
            candidates.append(
                SyndicatedCandidate(
                    title=title,
                    url=urljoin("https://www.brecorder.com/", href),
                    source="Business Recorder",
                )
            )
    except (urllib.error.URLError, TimeoutError, ValueError):
        pass
    _BUSINESS_RECORDER_CANDIDATES = candidates
    return candidates


def dawn_candidates() -> list[SyndicatedCandidate]:
    global _DAWN_CANDIDATES
    if _DAWN_CANDIDATES is not None:
        return _DAWN_CANDIDATES

    candidates: list[SyndicatedCandidate] = []
    try:
        markup = fetch_markup("https://www.dawn.com/latest-news/", timeout=18)
        for attrs, title in collect_anchors(markup):
            href = attrs.get("href", "")
            if len(title) < 20 or "/news/" not in href:
                continue
            candidates.append(
                SyndicatedCandidate(
                    title=title,
                    url=urljoin("https://www.dawn.com/", href),
                    source="Dawn",
                )
            )
    except (urllib.error.URLError, TimeoutError, ValueError):
        pass
    _DAWN_CANDIDATES = candidates
    return candidates


def aib_candidates() -> list[SyndicatedCandidate]:
    global _AIB_CANDIDATES
    if _AIB_CANDIDATES is not None:
        return _AIB_CANDIDATES

    page_urls = [
        "https://www.aib.ie/fxcentre/i-want-to/read-news",
        (
            "https://www.aib.ie/fxcentre/i-want-to/"
            "read-news.urncolonnewsmlcolonreuters_comcolon20260616colonnL1N42O03W"
        ),
    ]
    candidates: list[SyndicatedCandidate] = []
    seen: set[str] = set()

    def add_page(markup: str) -> None:
        for attrs, title in collect_anchors(markup):
            selector = attrs.get("data-id", "")
            if not selector.startswith("urncolonnewsmlcolonreuters_comcolon"):
                continue
            if selector in seen:
                continue
            seen.add(selector)
            candidates.append(
                SyndicatedCandidate(
                    title=title,
                    url=(
                        "https://www.aib.ie/content/aib/fxcentre/i-want-to/"
                        f"read-news.news-story.{selector}.html"
                    ),
                    source="AIB Reuters feed",
                )
            )

    for page_url in page_urls:
        try:
            add_page(fetch_markup(page_url, timeout=18))
        except (urllib.error.URLError, TimeoutError, ValueError):
            continue
    _AIB_CANDIDATES = candidates
    return candidates


def streetinsider_candidates() -> list[SyndicatedCandidate]:
    global _STREETINSIDER_CANDIDATES
    if _STREETINSIDER_CANDIDATES is not None:
        return _STREETINSIDER_CANDIDATES

    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar())
    )
    page_url = "https://www.streetinsider.com/Reuters/?classic=1"
    candidates: list[SyndicatedCandidate] = []
    seen: set[str] = set()

    for _ in range(10):
        try:
            markup = fetch_markup(page_url, timeout=18, opener=opener)
        except (urllib.error.URLError, TimeoutError, ValueError):
            break
        older_url = ""
        for attrs, title in collect_anchors(markup):
            href = html.unescape(attrs.get("href", ""))
            if title.lower() == "view older stories" and "before_id=" in href:
                older_url = urljoin("https://www.streetinsider.com/", href)
            if not href.startswith("Reuters/") or not href.endswith(".html"):
                continue
            article_url = urljoin("https://www.streetinsider.com/", href)
            if article_url in seen:
                continue
            seen.add(article_url)
            candidates.append(
                SyndicatedCandidate(
                    title=title,
                    url=article_url,
                    source="StreetInsider Reuters feed",
                )
            )
        if not older_url:
            break
        page_url = older_url + ("&classic=1" if "?" in older_url else "?classic=1")

    _STREETINSIDER_CANDIDATES = candidates
    return candidates


def recover_reuters_article(headline: str) -> tuple[str, str] | None:
    if not headline:
        return None

    providers = [
        business_recorder_candidates,
        aib_candidates,
        streetinsider_candidates,
    ]
    for provider in providers:
        candidates = provider()
        if not candidates:
            continue
        ranked = sorted(
            ((headline_score(headline, candidate.title), candidate) for candidate in candidates),
            key=lambda pair: pair[0],
            reverse=True,
        )
        if not ranked or ranked[0][0] < 0.84:
            continue
        candidate = ranked[0][1]
        try:
            recovered_text, recovered_title = extract_from_markup(
                fetch_markup(candidate.url, timeout=18)
            )
        except (urllib.error.URLError, TimeoutError, ValueError):
            continue
        if len(recovered_text) >= 200:
            return recovered_text, recovered_title or candidate.title
    return None


def recover_article_without_url(
    headline: str,
    source: str,
) -> tuple[str, str, str] | None:
    source_lower = source.lower()
    if "dawn" in source_lower:
        candidates = dawn_candidates()
    elif "business recorder" in source_lower:
        candidates = business_recorder_candidates()
    else:
        return None

    ranked = sorted(
        ((headline_score(headline, candidate.title), candidate) for candidate in candidates),
        key=lambda pair: pair[0],
        reverse=True,
    )
    if not ranked or ranked[0][0] < 0.84:
        return None
    candidate = ranked[0][1]
    try:
        recovered_text, recovered_title = extract_from_markup(
            fetch_markup(candidate.url, timeout=18)
        )
    except (urllib.error.URLError, TimeoutError, ValueError):
        return None
    if len(recovered_text) < 200:
        return None
    return recovered_text, recovered_title or candidate.title, candidate.url


def extract_article_from_url(
    url: str,
    headline: str = "",
    source: str = "",
) -> tuple[str, str, str]:
    if requires_ocr(url):
        return (
            "[This link contains an image or PDF scan. OCR or manually supplied article text is required.]",
            f"OCR is required for {url}.",
            "",
        )
    try:
        markup = fetch_markup(url)
    except (urllib.error.URLError, TimeoutError, ValueError) as exc:
        host = (urlparse(url).hostname or "").lower()
        if host == "reuters.com" or host.endswith(".reuters.com") or source.lower() == "reuters":
            recovered = recover_reuters_article(headline)
            if recovered:
                recovered_text, recovered_title = recovered
                return recovered_text, "", recovered_title
        return (
            "[The app could not extract this article automatically. Please paste the article text in the Excel template.]",
            f"URL extraction failed for {url}: {exc}",
            "",
        )

    text, extracted_title = extract_from_markup(markup)
    if len(text) < 120:
        host = (urlparse(url).hostname or "").lower()
        if host == "reuters.com" or host.endswith(".reuters.com") or source.lower() == "reuters":
            recovered = recover_reuters_article(headline)
            if recovered:
                recovered_text, recovered_title = recovered
                return recovered_text, "", recovered_title
        return (
            "[The app could not detect the article body. Please paste the article text in the Excel template.]",
            f"URL extraction returned too little text for {url}.",
            extracted_title,
        )
    return text, "", extracted_title


def is_boilerplate(text: str) -> bool:
    lowered = text.lower()
    blocked = [
        "subscribe",
        "advertisement",
        "all rights reserved",
        "sign in",
        "accept cookies",
        "privacy policy",
        "related stories",
        "share this",
    ]
    return any(term in lowered for term in blocked)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    r_pr.rFonts.set(qn("w:cs"), "Jameel Noori Nastaleeq")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para_spacing(paragraph, before=0, after=6, line=1.1):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    set_para_spacing(paragraph, after=0, line=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


class LinkBuilder:
    def __init__(self):
        self.bookmark_counter = 1

    def bookmark(self, paragraph, name: str) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_counter))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_counter))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)
        self.bookmark_counter += 1

    def internal_link(self, paragraph, text: str, anchor: str) -> None:
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_style = OxmlElement("w:rStyle")
        r_style.set(qn("w:val"), "Hyperlink")
        r_pr.append(r_style)
        run.append(r_pr)
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def external_link(self, paragraph, text: str, url: str) -> None:
        if not url:
            return
        r_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_style = OxmlElement("w:rStyle")
        r_style.set(qn("w:val"), "Hyperlink")
        r_pr.append(r_style)
        run.append(r_pr)
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)


def slug(text: str, fallback: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")
    return (cleaned[:36] or fallback)


def apply_core_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Jameel Noori Nastaleeq")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    style_tokens = [
        ("Heading 1", 15, BLUE, 14, 7),
        ("Heading 2", 12.5, BLUE, 10, 5),
        ("Heading 3", 11.5, NAVY, 7, 3),
    ]
    for style_name, size, color, before, after in style_tokens:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:cs"), "Jameel Noori Nastaleeq")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def set_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("Daily News Bulletin")
    set_run_font(run, size=8.5, color=GRAY, bold=True)
    run = header.add_run("    External Communications Department")
    set_run_font(run, size=8.5, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Generated by Daily Bulletin System MVP")
    set_run_font(run, size=8, color=GRAY)


def add_masthead(doc: Document, report_date: str, prepared_for: str, links: LinkBuilder, sections: list[str]) -> None:
    title = doc.add_paragraph()
    set_para_spacing(title, before=2, after=1)
    run = title.add_run("Daily News Bulletin")
    set_run_font(run, size=20, bold=True, color=DARK_BLUE)
    links.bookmark(title, "top")

    subtitle = doc.add_paragraph()
    set_para_spacing(subtitle, after=10)
    run = subtitle.add_run("External Communications Department")
    set_run_font(run, size=11, color=GRAY, bold=True)

    meta = doc.add_paragraph()
    set_para_spacing(meta, after=12)
    for label, value in [
        ("Date: ", report_date),
        ("Prepared for: ", prepared_for),
    ]:
        run = meta.add_run(label)
        set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
        run = meta.add_run(value + "    ")
        set_run_font(run, size=10.5, color=RGBColor(30, 41, 59))

    table = doc.add_table(rows=1, cols=len(sections))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, section_name in enumerate(sections):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "E8F2F0")
        set_cell_text(cell, section_name, bold=True, color=DARK_BLUE, size=8.8)
        link_para = cell.paragraphs[0]
        links.bookmark(link_para, f"nav_{index + 1}")


def section_order(items: Iterable[NewsItem]) -> list[str]:
    seen = list(DEFAULT_SECTIONS)
    for item in items:
        if item.section not in seen:
            seen.append(item.section)
    return seen


def group_items(items: list[NewsItem]) -> OrderedDict[str, OrderedDict[str, list[NewsItem]]]:
    grouped: OrderedDict[str, OrderedDict[str, list[NewsItem]]] = OrderedDict()
    for section in section_order(items):
        grouped[section] = OrderedDict()
    for item in items:
        grouped.setdefault(item.section, OrderedDict())
        grouped[item.section].setdefault(item.story_group, [])
        grouped[item.section][item.story_group].append(item)
    return grouped


def add_headline_index(doc: Document, grouped, links: LinkBuilder) -> dict[tuple[str, str], str]:
    anchors: dict[tuple[str, str], str] = {}
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.add_run("Headlines")
    links.bookmark(paragraph, "headlines")

    group_count = 1
    for section_name, stories in grouped.items():
        if not stories:
            continue
        heading = doc.add_paragraph(style="Heading 2")
        heading.add_run(section_name)
        links.bookmark(heading, f"section_{slug(section_name, 'section')}")

        for story_group, items in stories.items():
            anchor = f"article_{group_count}_{slug(story_group, 'story')}"
            anchors[(section_name, story_group)] = anchor
            group_count += 1

            p = doc.add_paragraph()
            set_para_spacing(p, after=2, line=1.05)
            p.paragraph_format.left_indent = Inches(0.16)
            links.internal_link(p, story_group, anchor)

            source_text = " / ".join(item.source for item in items)
            p = doc.add_paragraph()
            set_para_spacing(p, after=7, line=1.0)
            p.paragraph_format.left_indent = Inches(0.34)
            run = p.add_run(source_text)
            set_run_font(run, size=9.3, color=GRAY)

    return anchors


def add_separator(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    set_para_spacing(paragraph, after=8)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_article_body(doc: Document, grouped, anchors, links: LinkBuilder) -> None:
    doc.add_page_break()
    for section_name, stories in grouped.items():
        if not stories:
            continue
        section_heading = doc.add_paragraph(style="Heading 1")
        section_heading.add_run(section_name)

        for story_group, items in stories.items():
            story_heading = doc.add_paragraph(style="Heading 2")
            story_heading.add_run(story_group)
            links.bookmark(story_heading, anchors[(section_name, story_group)])

            for item in items:
                article_heading = doc.add_paragraph(style="Heading 3")
                article_heading.add_run(item.headline)
                if item.language.lower() == "urdu":
                    article_heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                source = doc.add_paragraph()
                set_para_spacing(source, after=5)
                run = source.add_run(source_meta(item))
                set_run_font(run, size=10, bold=True, color=DARK_BLUE)
                if item.priority:
                    run = source.add_run(f" | Priority: {item.priority}")
                    set_run_font(run, size=9.2, color=GRAY)
                if is_http_url(item.url):
                    run = source.add_run(" | ")
                    set_run_font(run, size=9.2, color=GRAY)
                    links.external_link(source, "Original URL", item.url)

                status_note = ""
                if item.status and item.status.lower() != "ok":
                    status_note = f"Extraction status: {item.status}"
                note_text = " | ".join(part for part in [item.notes, status_note] if part)
                if note_text:
                    notes = doc.add_paragraph()
                    set_para_spacing(notes, after=5)
                    run = notes.add_run(f"Note: {note_text}")
                    set_run_font(run, size=9.2, italic=True, color=GRAY)

                for body_paragraph in split_body(item.article_text):
                    p = doc.add_paragraph()
                    set_para_spacing(p, after=5, line=1.08)
                    if item.language.lower() == "urdu":
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    run = p.add_run(body_paragraph)
                    set_run_font(run, size=10.3, color=RGBColor(17, 24, 39))

                back = doc.add_paragraph()
                set_para_spacing(back, before=1, after=9)
                links.internal_link(back, "Back to Headlines", "headlines")

            add_separator(doc)


def split_body(text: str) -> list[str]:
    parts = [clean_text(part) for part in re.split(r"\n\s*\n", text or "")]
    return [part for part in parts if part] or ["[No article body available.]"]


def scrub_metadata(docx_path: Path) -> None:
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>Daily Bulletin System</dc:creator>", text)
                text = re.sub(
                    r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                    "<cp:lastModifiedBy>Daily Bulletin System</cp:lastModifiedBy>",
                    text,
                )
                data = text.encode("utf-8")
            if item.filename.endswith(".xml"):
                text = data.decode("utf-8", errors="ignore")
                text = re.sub(r'\s+w:rsid\w+="[^"]+"', "", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp_path.replace(docx_path)


def resolve_report_date(report_date: str | None) -> str:
    return report_date or date.today().strftime("%A, %B %d, %Y")


def filename_date() -> str:
    return date.today().strftime("%Y%m%d")


def build_metadata(items: list[NewsItem], grouped, warnings: list[str]) -> dict:
    return {
        "article_count": len(items),
        "story_count": sum(len(stories) for stories in grouped.values()),
        "section_count": sum(1 for stories in grouped.values() if stories),
        "warnings": warnings,
    }


def write_bulletin_docx(
    items: list[NewsItem],
    grouped,
    output_dir: str | Path,
    report_date: str,
    prepared_for: str,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    output_path = output_dir / f"Daily_News_Bulletin_{filename_date()}.docx"

    doc = Document()
    links = LinkBuilder()
    apply_core_styles(doc)
    set_header_footer(doc)
    add_masthead(doc, report_date, prepared_for, links, section_order(items))
    anchors = add_headline_index(doc, grouped, links)
    add_article_body(doc, grouped, anchors, links)
    doc.save(output_path)
    scrub_metadata(output_path)
    return output_path


def render_preview_html(items: list[NewsItem], grouped, report_date: str, prepared_for: str) -> str:
    def esc(value: str) -> str:
        return html.escape(value or "", quote=True)

    active_sections = [(section, stories) for section, stories in grouped.items() if stories]
    anchors: dict[tuple[str, str], str] = {}
    story_number = 1
    for section_name, stories in active_sections:
        for story_group in stories:
            anchors[(section_name, story_group)] = f"article-{story_number}-{slug(story_group, 'story').lower()}"
            story_number += 1

    parts = [
        "<!doctype html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8" />',
        '<meta name="viewport" content="width=device-width, initial-scale=1" />',
        "<title>Daily News Bulletin Preview</title>",
        "<style>",
        """
        :root {
          --ink: #111827;
          --muted: #5f6c70;
          --line: #d5dddf;
          --green: #1f6f68;
          --green-dark: #143d3a;
          --gold: #b98416;
          --paper: #ffffff;
          --page: #f3f6f5;
        }
        * { box-sizing: border-box; }
        body {
          margin: 0;
          background: var(--page);
          color: var(--ink);
          font-family: "Segoe UI", Arial, sans-serif;
          line-height: 1.58;
        }
        main {
          width: min(980px, calc(100vw - 28px));
          margin: 24px auto;
          background: var(--paper);
          border: 1px solid var(--line);
          border-radius: 8px;
          padding: 34px;
        }
        header {
          border-bottom: 4px solid var(--green);
          padding-bottom: 18px;
          margin-bottom: 24px;
        }
        h1, h2, h3, h4, p { letter-spacing: 0; }
        h1 {
          margin: 0;
          color: var(--green-dark);
          font-size: 34px;
          line-height: 1.1;
        }
        h2 {
          margin: 30px 0 12px;
          color: var(--green);
          font-size: 23px;
          border-bottom: 1px solid var(--line);
          padding-bottom: 7px;
        }
        h3 {
          margin: 22px 0 6px;
          color: var(--green-dark);
          font-size: 20px;
        }
        h4 {
          margin: 15px 0 5px;
          color: #17324d;
          font-size: 16px;
        }
        a { color: var(--green); font-weight: 700; }
        .subtitle, .meta, .source, .notes, .sources {
          color: var(--muted);
        }
        .subtitle {
          margin: 6px 0 12px;
          font-weight: 700;
        }
        .meta {
          display: flex;
          flex-wrap: wrap;
          gap: 16px;
          margin: 0;
          font-size: 14px;
        }
        .nav {
          display: flex;
          flex-wrap: wrap;
          gap: 8px;
          margin: 18px 0 0;
          padding: 0;
          list-style: none;
        }
        .nav a {
          display: inline-flex;
          border: 1px solid var(--line);
          border-radius: 999px;
          padding: 6px 10px;
          text-decoration: none;
          background: #f8fbfa;
          font-size: 13px;
        }
        .headlines {
          margin: 0 0 24px;
          padding-left: 22px;
        }
        .headlines li { margin: 11px 0; }
        .sources {
          display: block;
          margin-top: 2px;
          font-size: 13px;
          font-weight: 600;
        }
        .article {
          border-bottom: 1px solid var(--line);
          padding-bottom: 18px;
          margin-bottom: 18px;
        }
        .source, .notes {
          margin: 0 0 9px;
          font-size: 13px;
        }
        .notes { font-style: italic; }
        .body p { margin: 0 0 11px; }
        .urdu {
          direction: rtl;
          text-align: right;
          font-family: "Jameel Noori Nastaleeq", "Noto Nastaliq Urdu", serif;
        }
        .back {
          display: inline-flex;
          margin-top: 8px;
          font-size: 13px;
        }
        @media print {
          body { background: white; }
          main { width: auto; margin: 0; border: 0; padding: 0; }
          .nav, .back { display: none; }
        }
        @media (max-width: 640px) {
          main { padding: 22px; }
          h1 { font-size: 28px; }
        }
        """,
        "</style>",
        "</head>",
        "<body>",
        "<main>",
        '<header id="top">',
        "<h1>Daily News Bulletin</h1>",
        '<p class="subtitle">External Communications Department</p>',
        f'<p class="meta"><span><strong>Date:</strong> {esc(report_date)}</span><span><strong>Prepared for:</strong> {esc(prepared_for)}</span></p>',
        '<ul class="nav">',
    ]

    for section_name, _ in active_sections:
        parts.append(f'<li><a href="#section-{esc(slug(section_name, "section").lower())}">{esc(section_name)}</a></li>')
    parts.extend(["</ul>", "</header>", '<section id="headlines">', "<h2>Headlines</h2>"])

    for section_name, stories in active_sections:
        parts.append(f'<h3 id="section-{esc(slug(section_name, "section").lower())}">{esc(section_name)}</h3>')
        parts.append('<ol class="headlines">')
        for story_group, story_items in stories.items():
            source_text = " / ".join(source_meta(item) for item in story_items)
            anchor = anchors[(section_name, story_group)]
            parts.append(
                f'<li><a href="#{esc(anchor)}">{esc(story_group)}</a><span class="sources">{esc(source_text)}</span></li>'
            )
        parts.append("</ol>")

    parts.append("</section>")
    parts.append('<section id="articles">')
    for section_name, stories in active_sections:
        parts.append(f"<h2>{esc(section_name)}</h2>")
        for story_group, story_items in stories.items():
            anchor = anchors[(section_name, story_group)]
            parts.append(f'<article id="{esc(anchor)}" class="article">')
            parts.append(f"<h3>{esc(story_group)}</h3>")
            for item in story_items:
                article_class = ' class="urdu"' if item.language.lower() == "urdu" else ""
                parts.append(f"<h4{article_class}>{esc(item.headline)}</h4>")
                source_line = esc(source_meta(item))
                if item.priority:
                    source_line += f" | Priority: {esc(item.priority)}"
                if is_http_url(item.url):
                    source_line += f' | <a href="{esc(item.url)}" target="_blank" rel="noopener">Original URL</a>'
                parts.append(f'<p class="source">{source_line}</p>')
                status_note = ""
                if item.status and item.status.lower() != "ok":
                    status_note = f"Extraction status: {item.status}"
                note_text = " | ".join(part for part in [item.notes, status_note] if part)
                if note_text:
                    parts.append(f'<p class="notes">Note: {esc(note_text)}</p>')
                body_class = ' class="body urdu"' if item.language.lower() == "urdu" else ' class="body"'
                parts.append(f"<div{body_class}>")
                for paragraph in split_body(item.article_text):
                    parts.append(f"<p>{esc(paragraph)}</p>")
                parts.append("</div>")
            parts.append('<a class="back" href="#headlines">Back to Headlines</a>')
            parts.append("</article>")

    parts.extend(["</section>", "</main>", "</body>", "</html>"])
    return "\n".join(parts)


def write_preview_html(
    items: list[NewsItem],
    grouped,
    output_dir: str | Path,
    report_date: str,
    prepared_for: str,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"Daily_News_Bulletin_{filename_date()}_Preview.html"
    output_path.write_text(render_preview_html(items, grouped, report_date, prepared_for), encoding="utf-8")
    return output_path


def generate_bulletin(
    excel_path: str | Path,
    output_dir: str | Path,
    report_date: str | None = None,
    prepared_for: str = "SBP Management Forum",
) -> tuple[Path, dict]:
    excel_path = Path(excel_path)
    items, warnings = read_news_items(excel_path)
    grouped = group_items(items)
    resolved_date = resolve_report_date(report_date)
    output_path = write_bulletin_docx(items, grouped, output_dir, resolved_date, prepared_for)
    return output_path, build_metadata(items, grouped, warnings)


def generate_preview_html(
    excel_path: str | Path,
    output_dir: str | Path,
    report_date: str | None = None,
    prepared_for: str = "SBP Management Forum",
) -> tuple[Path, dict]:
    excel_path = Path(excel_path)
    items, warnings = read_news_items(excel_path)
    grouped = group_items(items)
    resolved_date = resolve_report_date(report_date)
    output_path = write_preview_html(items, grouped, output_dir, resolved_date, prepared_for)
    return output_path, build_metadata(items, grouped, warnings)


def generate_bulletin_artifacts(
    excel_path: str | Path,
    report_dir: str | Path,
    preview_dir: str | Path,
    report_date: str | None = None,
    prepared_for: str = "SBP Management Forum",
) -> tuple[Path, Path, dict]:
    excel_path = Path(excel_path)
    items, warnings = read_news_items(excel_path)
    grouped = group_items(items)
    resolved_date = resolve_report_date(report_date)
    report_path = write_bulletin_docx(items, grouped, report_dir, resolved_date, prepared_for)
    preview_path = write_preview_html(items, grouped, preview_dir, resolved_date, prepared_for)
    return report_path, preview_path, build_metadata(items, grouped, warnings)


def main(argv: list[str]) -> int:
    if len(argv) < 3:
        print("Usage: python bulletin_generator.py <input.xlsx> <output_dir> [report_date] [prepared_for]")
        return 2
    report_date = argv[3] if len(argv) >= 4 else None
    prepared_for = argv[4] if len(argv) >= 5 else "SBP Management Forum"
    output, metadata = generate_bulletin(argv[1], argv[2], report_date, prepared_for)
    print(output)
    print(metadata)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
